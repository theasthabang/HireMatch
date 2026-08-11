import os
import logging
from docx import Document
import pdfplumber

logger = logging.getLogger(__name__)

def extract_text(file_path: str, content_type: str) -> str:
    """
    Extracts and cleans text from a PDF or DOCX file.

    :param file_path: Path to the target file on disk.
    :param content_type: Content type or MIME type of the file.
    :return: Cleaned text string.
    :raises ValueError: For unsupported files, too large files, or extraction errors.
    """
    # 1. File size check (5MB limit)
    MAX_FILE_SIZE = 5 * 1024 * 1024  # 5MB in bytes
    try:
        file_size = os.path.getsize(file_path)
    except Exception as e:
        logger.error(f"Error accessing file path '{file_path}': {e}")
        raise ValueError(f"Could not access the file for reading. Error: {str(e)}")

    if file_size > MAX_FILE_SIZE:
        logger.error(f"File size exceeds 5MB limit. Size: {file_size} bytes.")
        raise ValueError("File is too large. The maximum allowed size is 5MB.")

    # 2. File type detection
    content_type_lower = content_type.lower()
    _, ext = os.path.splitext(file_path.lower())

    is_pdf = "pdf" in content_type_lower or ext == ".pdf"
    is_docx = (
        "wordprocessingml.document" in content_type_lower
        or ext == ".docx"
        or "docx" in content_type_lower
    )

    if not (is_pdf or is_docx):
        detected_format = ext or content_type
        logger.error(f"Unsupported file format: content_type={content_type}, extension={ext}")
        raise ValueError(
            f"Unsupported file format '{detected_format}'. Please upload a PDF or DOCX file."
        )

    raw_text_parts = []

    # 3. Extraction block
    if is_pdf:
        logger.info(f"Detected PDF file type for '{file_path}'")
        try:
            with pdfplumber.open(file_path) as pdf:
                pages_processed = 0
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        raw_text_parts.append(page_text)
                        pages_processed += 1
                logger.info(f"Successfully processed {pages_processed} pages from PDF.")
        except Exception as e:
            logger.error(f"Failed to extract text from PDF '{file_path}': {e}")
            raise ValueError(
                f"Failed to parse PDF file. The file may be corrupt or encrypted. Details: {str(e)}"
            )

    elif is_docx:
        logger.info(f"Detected DOCX file type for '{file_path}'")
        try:
            doc = Document(file_path)
            
            # Extract paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    raw_text_parts.append(paragraph.text)
            
            # Extract table cells (since resumes often format work experience inside tables)
            table_cells_extracted = 0
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            raw_text_parts.append(cell.text)
                            table_cells_extracted += 1
            
            logger.info(
                f"Successfully processed {len(doc.paragraphs)} paragraphs and {table_cells_extracted} table cells from DOCX."
            )
        except Exception as e:
            logger.error(f"Failed to extract text from DOCX '{file_path}': {e}")
            raise ValueError(
                f"Failed to parse DOCX file. The file may be corrupt. Details: {str(e)}"
            )

    # 4. Clean whitespace and excessive blank lines
    combined_text = "\n".join(raw_text_parts)
    
    # Strip leading/trailing whitespaces, remove empty lines, collapse multiple spaces
    lines = [line.strip() for line in combined_text.splitlines()]
    cleaned_lines = [" ".join(line.split()) for line in lines if line]
    cleaned_text = "\n".join(cleaned_lines)
    
    char_count = len(cleaned_text)
    logger.info(f"Completed extraction. Cleaned character count: {char_count}")

    # 5. Length verification
    if char_count < 50:
        logger.error(f"Extracted content is too short ({char_count} characters) for '{file_path}'")
        raise ValueError(
            "Could not extract readable text from this file. Please upload a text-based PDF or DOCX."
        )

    return cleaned_text


def assess_parse_confidence(cleaned_text: str) -> tuple[str, str]:
    """
    Heuristically estimates how reliable the ATS score is likely to be, based
    on properties of the *already-cleaned* extracted text.

    This is intentionally NOT an LLM call — it's a handful of cheap,
    deterministic signals (line-length distribution, alphabetic density,
    fragment ratio) that tend to show up when a PDF used a multi-column or
    table-based layout, or when a scanned/image PDF produced mostly noise.
    A low-confidence signal here means "trust this score less", independent
    of whatever number the ATS chain produced.

    :param cleaned_text: The cleaned text returned by extract_text().
    :return: (confidence, reason) where confidence is "high" | "medium" | "low".
    """
    lines = [line for line in cleaned_text.splitlines() if line.strip()]
    if not lines:
        return "low", "No readable text lines were found in this document."

    line_lengths = [len(line) for line in lines]
    avg_len = sum(line_lengths) / len(line_lengths)

    # Fragment ratio: very short lines (<=3 chars) are a strong signal of a
    # multi-column or table layout being flattened incorrectly during
    # extraction (e.g. a date, a bullet glyph, and a single word each landing
    # on their own line instead of flowing together).
    short_line_count = sum(1 for l in line_lengths if l <= 3)
    fragment_ratio = short_line_count / len(lines)

    # Alphabetic density across the whole document — scanned/OCR-garbled PDFs
    # often extract as a high proportion of symbols/whitespace noise.
    alpha_chars = sum(1 for c in cleaned_text if c.isalpha())
    alpha_ratio = alpha_chars / max(1, len(cleaned_text))

    issues = []
    if fragment_ratio > 0.35:
        issues.append("a high proportion of very short, fragmented lines (possible table/column layout)")
    if avg_len < 15:
        issues.append("unusually short average line length")
    if alpha_ratio < 0.5:
        issues.append("a low ratio of readable alphabetic text to total characters")

    if len(issues) >= 2:
        return "low", "Detected " + " and ".join(issues) + "."
    if len(issues) == 1:
        return "medium", "Detected " + issues[0] + "."
    return "high", "Text extracted cleanly with no signs of layout or OCR issues."